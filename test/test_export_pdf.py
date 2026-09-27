from unittest.mock import MagicMock

import openpyxl as px
import pytest
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage

import src.process.export_pdf as export_pdf


def _make_png(path, width=1000, height=500):
    """テスト用のダミーpng画像を生成"""
    PILImage.new("RGB", (width, height), color="white").save(path)


# ----- _resize_image -----


def test_resize_image_keeps_aspect_ratio_and_sets_target_width(tmp_path):
    """縦横比を保ったまま幅がtarget_widthにリサイズされること"""
    png_path = tmp_path / "sample.png"
    _make_png(png_path, width=1000, height=500)  # 縦横比 1:0.5

    img = XLImage(str(png_path))
    original_ratio = img.height / img.width

    export_pdf._resize_image(img, target_width=650)

    # A4印刷可能領域に合わせた画像幅であること
    assert img.width == 650

    # 画像の高さがリサイズ前の縦横比×新しい幅(650)と一致していること
    assert img.height == pytest.approx(650 * original_ratio, abs=1)


# ----- _set_print_setup -----


def test_set_print_setup_applies_expected_settings():
    """印刷設定（A4・縦向き・余白・1ページ納め・中央配置）が適用されていること"""
    wb = px.Workbook()
    ws = wb.active

    # Pylance対策: Noneではないことを保証する
    assert ws is not None

    export_pdf._set_print_setup(ws)

    # 用紙サイズがA4サイズであること
    # openpyxlは代入時、paperSize定数(文字列)を内部でint型に変換して保持するため、intに変換して比較
    assert ws.page_setup.paperSize == int(ws.PAPERSIZE_A4)

    # 用紙の向きが縦向きであること
    assert ws.page_setup.orientation == ws.ORIENTATION_PORTRAIT

    # 上方向の余白が1.0であること
    assert ws.page_margins.top == 1.0

    # 下方向の余白が1.0であること
    assert ws.page_margins.bottom == 1.0

    # 右方向の余白が0.75であること
    assert ws.page_margins.right == 0.75

    # 左方向の余白が0.75であること
    assert ws.page_margins.left == 0.75

    # Pylance対策
    assert ws.sheet_properties.pageSetUpPr is not None

    # A4用紙1ページに収まっていること
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True

    # 用紙の横部分が1ページに収められていること
    assert ws.page_setup.fitToWidth == 1

    # 用紙の縦部分が1ページを超えた場合に、次のページで表示されていること
    assert ws.page_setup.fitToHeight == 0

    # ページ中央が水平に設定されていること
    assert ws.print_options.horizontalCentered is True


# ----- plot_png_paste -----


class FakeFolderPath:
    """FolderPahtの代替として、テスト用の一時フォルダを返す"""

    def __init__(self, excel_paths, pie_dir: str, bar_dir: str):
        self._excel_paths = excel_paths
        self._pie_dir = pie_dir
        self._bar_dir = bar_dir

    def get_excel_path(self):
        return self._excel_paths

    def get_plot_path(self, plot_type):
        return self._pie_dir if plot_type == "pie_plot" else self._bar_dir


@pytest.fixture
def excel_path_with_plots(tmp_path, monkeypatch):
    """加工済みExcelファイルとダミーpng画像を用意し、FolderPathを差し替える"""
    # tmp_path = pytest用の仮ディレクトリ
    processed_dir = tmp_path / "processed"
    processed_dir.mkdir()

    pie_dir = tmp_path / "pie_plot"
    pie_dir.mkdir()

    bar_dir = tmp_path / "bar_plot"
    bar_dir.mkdir()

    excel_path = processed_dir / "Sheet1.xlsx"

    wb = px.Workbook()
    ws = wb.active

    # Pylance対策: None出ないことを保証する
    assert ws is not None

    ws.title = "Sheet1"
    wb.save(excel_path)

    _make_png(pie_dir / "Sheet1売上円グラフ.png")
    _make_png(bar_dir / "Sheet1売上棒グラフ.png")

    # export_pdf.py内のFolderPathクラスを差し替え（元のpath_general.pyではない）
    monkeypatch.setattr(
        export_pdf,
        "FolderPath",
        lambda: FakeFolderPath(
            # 複数Excelファイルを扱う処理にしているため、excel_pathsをリストとして渡す
            excel_paths=[excel_path], pie_dir=pie_dir, bar_dir=bar_dir
        ),
    )

    return excel_path


def test_plot_png_paste_attaches_images_to_graph_sheet(excel_path_with_plots):
    """円グラフ・棒グラフpngが加工済みExcelファイルの「グラフ」シートに添付されていること"""
    export_pdf.plot_png_paste()

    wb = px.load_workbook(excel_path_with_plots)

    # workbook内に「グラフ」シートが存在していること
    assert "グラフ" in wb.sheetnames
    ws = wb["グラフ"]

    # Pylance対策: wsのインスタンスがWorksheetであることを保証する
    assert isinstance(ws, Worksheet)

    # 画像が2枚（円グラフ・棒グラフ）が添付されていること
    assert len(ws._images) == 2  # type: ignore[attr-defined]


def test_plot_png_paste_images_placed_at_expected_anchors(excel_path_with_plots):
    """画像がA1・A30セルを起点に配置されていること"""
    export_pdf.plot_png_paste()

    wb = px.load_workbook(excel_path_with_plots)
    ws = wb["グラフ"]

    # Pylance対策: wsのインスタンスがWorksheetであることを保証する
    assert isinstance(ws, Worksheet)

    # アンカー行番号は0始まりのため、A1 = 0,  A30 = 29
    anchor_rows = sorted(img.anchor._from.row for img in ws._images)  # type: ignore[attr-defined]

    # 画像がA1、A30セルを起点に配置されていること
    assert anchor_rows == [0, 29]


def test_plot_png_paste_raises_when_pie_png_missing(tmp_path, monkeypatch):
    """円グラフpngが存在しない場合にFileNotFoundErrorが送出されること"""
    # tmp_path = pytest用の仮ディレクトリ
    processed_dir = tmp_path / "processed"
    processed_dir.mkdir()

    pie_dir = tmp_path / "pie_plot"
    pie_dir.mkdir()

    bar_dir = tmp_path / "bar_plot"
    bar_dir.mkdir()

    excel_path = processed_dir / "Sheet1.xlsx"

    wb = px.Workbook()
    ws = wb.active

    # Pylance対策: Noneではないことを保証する
    assert ws is not None

    ws.title = "Sheet1"
    wb.save(excel_path)

    # 棒グラフpngのみ用意し、円グラフpngは用意しない
    _make_png(bar_dir / "Sheet1売上棒グラフ.png")

    # export_pdf.py内のFolderPathクラスを差し替え（元のpath_general.pyではない）
    monkeypatch.setattr(
        export_pdf,
        "FolderPath",
        lambda: FakeFolderPath(
            # 複数Excelファイルを扱う処理にしているため、excel_pathsをリストとして渡す
            excel_paths=[excel_path], pie_dir=pie_dir, bar_dir=bar_dir
        ),
    )

    with pytest.raises(FileNotFoundError, match="円グラフpng"):
        export_pdf.plot_png_paste()


def test_plot_png_paste_raises_when_bar_png_missing(tmp_path, monkeypatch):
    """棒グラフpngが存在しない場合にFileNotFoundErrorが送出されること"""
    # tmp_path = pytest用の仮ディレクトリ
    processed_dir = tmp_path / "processed"
    processed_dir.mkdir()

    pie_dir = tmp_path / "pie_plot"
    pie_dir.mkdir()

    bar_dir = tmp_path / "bar_plot"
    bar_dir.mkdir()

    excel_path = processed_dir / "Sheet1.xlsx"

    wb = px.Workbook()
    ws = wb.active

    # Pylance対策: Noneではないことを保証する
    assert ws is not None

    ws.title = "Sheet1"
    wb.save(excel_path)

    # 円グラフpngのみ用意し、棒グラフpngは用意しない
    _make_png(pie_dir / "Sheet1売上円グラフ.png")

    # export_pdf.py内のFolderPathクラスを差し替え（元のpath_general.pyではない）
    monkeypatch.setattr(
        export_pdf,
        "FolderPath",
        lambda: FakeFolderPath(
            # 複数Excelファイルを扱う処理にしているため、excel_pathsをリストとして渡す
            excel_paths=[excel_path], pie_dir=pie_dir, bar_dir=bar_dir
        ),
    )

    with pytest.raises(FileNotFoundError, match="棒グラフpng"):
        export_pdf.plot_png_paste()


# ----- create_pdf -----


class FakeFolderPathForPdf:
    """FolderPathの代替として、テスト用の一時フォルダを返す"""

    def __init__(self, excel_paths, pdf_dir):
        self._excel_paths = excel_paths
        self._pdf_dir = pdf_dir

    def get_excel_path(self):
        return self._excel_paths

    def get_pdf_path(self):
        return self._pdf_dir


def test_create_pdf_exports_excel_as_pdf(tmp_path, monkeypatch):
    """加工済みExcelファイルがPDFとして出力されていること"""
    excel_path = tmp_path / "Sheet1.xlsx"
    excel_path.touch()

    pdf_dir = tmp_path / "pdf"
    pdf_dir.mkdir()

    # 本物のExcelアプリの代わりに、なんでも受け止めるダミーを用意
    fake_excel_app = MagicMock()
    fake_wb = fake_excel_app.Workbooks.Open.return_value
    fake_wb.ActiveSheet = fake_wb  # Workbook.Select()後、ActiveSheet自身とみなす

    # export_pdf.py内のFolderPathクラスを差し替え（元のpath_general.pyではない）
    monkeypatch.setattr(
        export_pdf,
        "FolderPath",
        # 複数Excelファイルを扱う処理にしているため、excel_pathsをリストとして渡す
        lambda: FakeFolderPathForPdf(excel_paths=[excel_path], pdf_dir=pdf_dir),
    )

    monkeypatch.setattr(export_pdf.win32, "EnsureDispatch", lambda name: fake_excel_app)

    export_pdf.create_pdf()

    expected_pdf_path = str(pdf_dir / "Sheet1.pdf")

    # 0 = pdf形式
    # 加工済みExcelファイルがPDFとして出力されていること
    fake_wb.ActiveSheet.ExportAsFixedFormat.assert_called_once_with(
        0, expected_pdf_path
    )


def test_create_pdf_closes_workbook_without_saveing_and_quits_excel(
    tmp_path, monkeypatch
):
    """PDF出力後、変更を保存せずワークブックを閉じ、Excelを終了していること"""
    excel_path = tmp_path / "Sheet1.xlsx"
    excel_path.touch()

    pdf_dir = tmp_path / "pdf"
    pdf_dir.mkdir()

    fake_excel_app = MagicMock()
    fake_wb = fake_excel_app.Workbooks.Open.return_value
    fake_wb.ActiveSheet = fake_wb

    # export_pdf.py内のFolderPathクラスを差し替え（元のpath_general.pyではない）
    monkeypatch.setattr(
        export_pdf,
        "FolderPath",
        # 複数Excelファイルを扱う処理にしているため、excel_pathsをリストとして渡す
        lambda: FakeFolderPathForPdf(excel_paths=[excel_path], pdf_dir=pdf_dir),
    )

    # export_pdf.py内のwin32.EnsureDispatch("Excel.Application")を差し替え
    monkeypatch.setattr(export_pdf.win32, "EnsureDispatch", lambda name: fake_excel_app)

    export_pdf.create_pdf()

    # 変更を保存せずにワークブックを閉じていること
    fake_wb.Close.assert_called_once_with(SaveChanges=False)
    
    # Excelを終了していること
    fake_excel_app.Quit.assert_called_once()
