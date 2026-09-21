import sys
from pathlib import Path

import openpyxl as px
import win32com.client.gencache as win32
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.worksheet import Worksheet  # 型ヒント記述用

from src.common.path_general import FolderPath


def _resize_image(img: XLImage, target_width: int) -> None:
    """画像の縦横比を保ったまま幅を基準にリサイズ"""
    ratio = target_width / img.width
    img.width = target_width
    img.height = int(img.height * ratio)


def _set_print_setup(ws: Worksheet) -> None:
    """印刷設定（A4・縦向き・余白・1ページ納め）を統一"""
    # A4・縦向き・余白設定
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_margins.top = 1.0
    ws.page_margins.bottom = 1.0
    ws.page_margins.right = 0.75
    ws.page_margins.left = 0.75

    assert ws.sheet_properties.pageSetUpPr is not None
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 横は1ページに納める、縦の場合は1ページを超える場合は次のページで表示
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # 0 = 高さ制限なし

    ws.print_options.horizontalCentered = True


def plot_png_paste():
    """加工済みExcelファイル内にグラフ結果pngファイルを添付"""
    folder_path = FolderPath()

    excel_paths = folder_path.get_excel_path()

    pie_plot_path = folder_path.get_plot_path("pie_plot")
    bar_plot_path = folder_path.get_plot_path("bar_plot")

    # A4印刷可能領域に合わせた画像幅（px、余白19mm・96dpi換算の目安）
    IMAGE_WIDTH = 650

    for excel_path in excel_paths:
        wb = px.load_workbook(excel_path)

        # 既存シート名を基準にpngファイル名を組み立て
        data_ws = wb.worksheets[0]
        sheet_name = data_ws.title
        _set_print_setup(data_ws)

        pie_png_path = pie_plot_path / f"{sheet_name}売上円グラフ.png"
        bar_png_path = bar_plot_path / f"{sheet_name}売上棒グラフ.png"

        if not pie_png_path.exists():
            raise FileNotFoundError(f"円グラフpngが見つかりません: {pie_png_path}")
        if not bar_png_path.exists():
            raise FileNotFoundError(f"棒グラフpngが見つかりません: {bar_png_path}")

        # グラフ貼り付け用の新規ワークシート追加（既に存在している場合は上書き）
        GRAPH = "グラフ"

        if GRAPH in wb.sheetnames:
            del wb[GRAPH]

        ws = wb.create_sheet(GRAPH)
        _set_print_setup(ws)

        pie_img = XLImage(str(pie_png_path))
        _resize_image(pie_img, IMAGE_WIDTH)

        bar_img = XLImage(str(bar_png_path))
        _resize_image(bar_img, IMAGE_WIDTH)

        ws.add_image(pie_img, "A1")
        ws.add_image(bar_img, "A30")

        wb.save(excel_path)
        print(f"pngファイルを添付しました: {excel_path.name}")


def create_pdf() -> None:
    """加工済ExcelファイルをPDFとして出力"""
    folder_path = FolderPath()

    excel_paths = folder_path.get_excel_path()
    pdf_dir = folder_path.get_pdf_path()

    excel = win32.EnsureDispatch("Excel.Application")
    excel.Visible = False

    try:
        for excel_path in excel_paths:
            wb = excel.Workbooks.Open(str(excel_path))
            try:
                wb.Worksheets.Select()  # 全シート選択
                pdf_path = pdf_dir / f"{excel_path.stem}.pdf"
                wb.ActiveSheet.ExportAsFixedFormat(0, str(pdf_path))
                print(f"PDF出力しました: {pdf_path.name}")
            finally:
                wb.Close(SaveChanges=False)
    finally:
        excel.Quit()


def main() -> None:
    try:
        plot_png_paste()
        create_pdf()
    except (FileNotFoundError, ValueError, KeyError, RuntimeError) as e:
        print(f"処理を中断しました: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"想定外のエラーが発生しました: {e}")
        sys.exit(1)
    finally:
        print("処理を実行しました")


if __name__ == "__main__":
    main()
