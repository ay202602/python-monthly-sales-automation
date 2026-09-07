import pandas as pd
import openpyxl as px
import unicodedata
from openpyxl.styles import PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from pathlib import Path

def get_file_path() -> Path:
    """データファイル格納パスを取得"""
    file_path = Path(__file__).resolve().parents[1] / "data"

    if not file_path.exists():
        raise FileNotFoundError("dataフォルダが存在しません。")
    
    return file_path

def display_width(text: str) ->int:
    """全角文字は2、半角文字は1としてカウントした表示幅を返す"""
    width = 0
    for ch in text:
        # F(fullwidth), W(Wide)は全角扱い、それ以外は半角扱い
        if unicodedata.east_asian_width(ch) in ("F", "W"):
            width += 2
        else:
            width += 1
    return width


def cell_decoration(processed: dict[str, pd.DataFrame]) -> dict[str, px.Workbook]:
    """加工済みのデータをExcelに装飾"""
    # DDEBF7 = 薄い青色
    header_fill = PatternFill(
        fill_type="solid", start_color="DDEBF7", end_color="DDEBF7"
    )
    # 格子罫線
    thin = Side(style="thin", color="000000")
    grid_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    workbooks: dict[str, px.Workbook] = {}

    for sheet_name, df in processed.items():
        wb = px.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = sheet_name

        for row in dataframe_to_rows(df, index=False, header=True):
            ws.append(row)

        # ws[1] = セル1行目全体
        for cell in ws[1]:
            cell.fill = header_fill

        for row in ws.iter_rows(
            min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column
        ):
            for cell in row:
                cell.border = grid_border
        
        # 会計の書式設定（Excelから引用）
        ACCOUNTING_FORMAT = '_ ¥* #,##0_ ;_ ¥* -#,##0_ ;_ ¥* "-"_ ;_ @_ '
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = ACCOUNTING_FORMAT
        
        for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
            max_length = 0
            for cell in col_cells:
                if cell.value is None:
                    continue
                if isinstance(cell.value, (int,float)):
                    display_text = f"{cell.value:,.0f}"
                else:
                    display_text = str(cell.value)
                max_length = max(max_length, display_width(display_text))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_length + 2
        
        workbooks[sheet_name] = wb

    return workbooks


def save_excel(workbooks: dict[str, px.Workbook]) -> None:
    """装飾したExcelファイルを出力"""
    PROCESSED_PATH = FILE_PATH / "processed"

    for sheet_name, wb in workbooks.items():
        wb.save(PROCESSED_PATH / f"{sheet_name}売上データ.xlsx")
        print(
            f"xlsxファイルの出力が完了しました。ファイル名：{sheet_name}売上データ.xlsx"
        )


def main() -> None:
    processed = data_process()
    workbooks = cell_decoration(processed)
    save_excel(workbooks)


if __name__ == "__main__":
    main()
